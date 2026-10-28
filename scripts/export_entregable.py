#!/usr/bin/env python3
"""Convierte el documento ground truth en Markdown a DOCX y PDF con el formato de entrega."""

import argparse
import os
import re
import subprocess
import sys
import venv
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent


def _venv_python(venv_dir: Path) -> Path:
    if os.name == "nt":
        return venv_dir / "Scripts" / "python.exe"
    return venv_dir / "bin" / "python"


def _ensure_dependencies():
    try:
        import docx  # noqa: F401
        import reportlab  # noqa: F401
        return
    except ImportError:
        pass

    if os.environ.get("SDG_VENV_BOOTSTRAPPED") == "1":
        raise RuntimeError(
            "No se pudieron importar python-docx y reportlab desde el entorno virtual."
        )

    venv_dir = REPO_ROOT / ".venv"
    vpy = _venv_python(venv_dir)
    req_file = REPO_ROOT / "requirements.txt"

    if not vpy.exists():
        print(f"Creando entorno virtual en {venv_dir}...")
        venv.create(str(venv_dir), with_pip=True)

    print("Instalando dependencias desde requirements.txt...")
    subprocess.check_call([str(vpy), "-m", "pip", "install", "-r", str(req_file)])

    env = os.environ.copy()
    env["SDG_VENV_BOOTSTRAPPED"] = "1"
    res = subprocess.run([str(vpy), str(Path(__file__).resolve()), *sys.argv[1:]], env=env)
    sys.exit(res.returncode)


_ensure_dependencies()

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


TOKEN_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def parse_inline_tokens(text: str):
    """Divide texto Markdown en tuplas (texto, bold, italic)."""
    tokens = []
    pos = 0
    for m in TOKEN_RE.finditer(text):
        if m.start() > pos:
            tokens.append((text[pos : m.start()], False, False))
        raw = m.group(0)
        if raw.startswith("**") and raw.endswith("**"):
            tokens.append((raw[2:-2], True, False))
        elif raw.startswith("*") and raw.endswith("*"):
            tokens.append((raw[1:-1], False, True))
        pos = m.end()
    if pos < len(text):
        tokens.append((text[pos:], False, False))
    return tokens


def inline_md_to_rl_xml(text: str, bold_color: str = "#1E1E1E") -> str:
    """Convierte inline Markdown a etiquetas XML válidas para ReportLab."""
    out = []
    for chunk, is_bold, is_italic in parse_inline_tokens(text):
        escaped = (
            chunk.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )
        if is_bold:
            out.append(f'<font color="{bold_color}"><b>{escaped}</b></font>')
        elif is_italic:
            out.append(f"<i>{escaped}</i>")
        else:
            out.append(escaped)
    return "".join(out)


def parse_markdown_blocks(md_text: str):
    """Parsea el archivo Markdown en bloques estructurados."""
    lines = md_text.splitlines()
    blocks = []
    i = 0
    n = len(lines)
    seen_h1 = False

    while i < n:
        line = lines[i].rstrip()

        if not line.strip():
            i += 1
            continue

        # Bloque de encabezado institucional previo al título principal (# )
        if not seen_h1 and line.startswith("**") and line.endswith("**"):
            header_lines = []
            while i < n and lines[i].strip().startswith("**") and lines[i].strip().endswith("**"):
                header_lines.append(lines[i].strip()[2:-2])
                i += 1
            blocks.append(("inst_header", header_lines))
            continue

        # Título principal (# )
        if line.startswith("# "):
            seen_h1 = True
            blocks.append(("title", line[2:].strip()))
            i += 1
            # Subtítulo en cursiva inmediatamente posterior
            while i < n and not lines[i].strip():
                i += 1
            sub_lines = []
            while (
                i < n
                and lines[i].strip().startswith("*")
                and not lines[i].strip().startswith("**")
                and lines[i].strip().endswith("*")
            ):
                sub_lines.append(lines[i].strip()[1:-1])
                i += 1
            if sub_lines:
                blocks.append(("subtitle", sub_lines))
            continue

        # Callout / Resumen ejecutivo (> )
        if line.startswith(">"):
            q_lines = []
            while i < n and lines[i].strip().startswith(">"):
                content = lines[i].strip()[1:].strip()
                if content:
                    q_lines.append(content)
                i += 1
            title_line = ""
            body_lines = []
            if q_lines and q_lines[0].startswith("**") and q_lines[0].endswith("**"):
                title_line = q_lines[0][2:-2]
                rest = q_lines[1:]
            else:
                rest = q_lines
            for rl in rest:
                if rl.startswith("*") and not rl.startswith("**") and rl.endswith("*"):
                    body_lines.append(rl[1:-1])
                else:
                    body_lines.append(rl)
            blocks.append(("callout", (title_line, " ".join(body_lines))))
            continue

        # Encabezados de sección
        if line.startswith("#### "):
            blocks.append(("h3", line[5:].strip()))
            i += 1
            continue
        if line.startswith("### "):
            blocks.append(("h2", line[4:].strip()))
            i += 1
            continue
        if line.startswith("## "):
            blocks.append(("h1", line[3:].strip()))
            i += 1
            continue

        # Tablas Markdown
        if line.startswith("|") and line.endswith("|"):
            tbl_rows = []
            while i < n and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                raw_row = lines[i].strip()
                cells = [c.strip() for c in raw_row.strip("|").split("|")]
                # Ignorar fila separadora | --- | --- |
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    tbl_rows.append(cells)
                i += 1
            if tbl_rows:
                blocks.append(("table", tbl_rows))
            continue

        # Ítems de lista (- )
        if line.startswith("- "):
            items = []
            while i < n and lines[i].strip().startswith("- "):
                items.append(lines[i].strip()[2:].strip())
                i += 1
            blocks.append(("bullets", items))
            continue

        # Párrafo normal
        para_lines = []
        while (
            i < n
            and lines[i].strip()
            and not lines[i].strip().startswith(("#", ">", "|", "- "))
        ):
            para_lines.append(lines[i].strip())
            i += 1
        blocks.append(("paragraph", " ".join(para_lines)))

    return blocks


def set_cell_shading(cell, fill_hex: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def set_cell_margins(cell, top: int, bottom: int, left: int, right: int):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_width_dxa(cell, width_dxa: int):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:type"), "dxa")
    tcW.set(qn("w:w"), str(width_dxa))


def add_Formatted_runs(p, text: str):
    tokens = parse_inline_tokens(text)
    for idx, (chunk, is_bold, is_italic) in enumerate(tokens):
        # Si un span en negrita va seguido de texto normal que empieza con espacio,
        # preservamos el espacio dentro del run en negrita igual que el DOCX original.
        if is_bold and idx + 1 < len(tokens):
            next_chunk, next_b, next_i = tokens[idx + 1]
            if not next_b and not next_i and next_chunk.startswith(" "):
                chunk = chunk + " "
                tokens[idx + 1] = (next_chunk[1:], next_b, next_i)
        if not chunk:
            continue
        r = p.add_run(chunk)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        if is_bold:
            r.bold = True
            r.font.color.rgb = RGBColor(0x1E, 0x1E, 0x1E)
        else:
            if is_italic:
                r.italic = True
            else:
                r.italic = False
            r.font.color.rgb = RGBColor(0x28, 0x28, 0x28)


def build_docx(blocks, out_path: Path):
    doc = Document()

    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    for btype, data in blocks:
        if btype == "inst_header":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(24)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run()
            r.font.name = "Arial"
            r.font.size = Pt(12)
            r.bold = True
            r.font.color.rgb = RGBColor(0x64, 0x64, 0x64)
            for idx, line in enumerate(data):
                if idx > 0:
                    r.add_break()
                r.add_text(line)

        elif btype == "title":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(18)
            r = p.add_run(data)
            r.font.name = "Arial"
            r.font.size = Pt(18)
            r.bold = True
            r.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

        elif btype == "subtitle":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(24)
            r = p.add_run()
            r.font.name = "Calibri"
            r.font.size = Pt(11)
            r.italic = True
            for idx, line in enumerate(data):
                if idx > 0:
                    r.add_break()
                r.add_text(line)

        elif btype == "callout":
            title_line, body_text = data
            spacer_before = doc.add_paragraph()
            spacer_before.paragraph_format.space_after = Pt(12)

            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            set_cell_width_dxa(cell, 9360)
            set_cell_shading(cell, "F0F4F8")
            set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = OxmlElement("w:tcBorders")
            left_b = OxmlElement("w:left")
            left_b.set(qn("w:val"), "single")
            left_b.set(qn("w:sz"), "24")
            left_b.set(qn("w:space"), "0")
            left_b.set(qn("w:color"), "003366")
            tcBorders.append(left_b)
            for side in ("top", "right", "bottom"):
                b = OxmlElement(f"w:{side}")
                b.set(qn("w:val"), "none")
                tcBorders.append(b)
            tcPr.append(tcBorders)

            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(0)
            cp.paragraph_format.space_after = Pt(2)
            if title_line:
                rt = cp.add_run(title_line)
                rt.font.name = "Arial"
                rt.font.size = Pt(11)
                rt.bold = True
                rt.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
                rt.add_break()
            rb = cp.add_run(body_text)
            rb.font.name = "Calibri"
            rb.font.size = Pt(10.5)
            rb.italic = True
            rb.font.color.rgb = RGBColor(0x32, 0x32, 0x32)

            spacer_after = doc.add_paragraph()
            spacer_after.paragraph_format.space_after = Pt(4)

        elif btype == "h1":
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(data)
            r.font.name = "Arial"
            r.font.size = Pt(16)
            r.bold = True
            r.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

        elif btype == "h2":
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(data)
            r.font.name = "Arial"
            r.font.size = Pt(13)
            r.bold = True
            r.font.color.rgb = RGBColor(0x29, 0x80, 0xB9)

        elif btype == "h3":
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(data)
            r.font.name = "Arial"
            r.font.size = Pt(11)
            r.bold = True
            r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        elif btype == "paragraph":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            add_Formatted_runs(p, data)

        elif btype == "bullets":
            for item in data:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.15
                add_Formatted_runs(p, item)

        elif btype == "table":
            num_rows = len(data)
            num_cols = max(len(r) for r in data)
            if num_cols == 2:
                col_widths_dxa = [2430, 6930]
            elif num_cols == 5:
                col_widths_dxa = [2060, 2250, 1220, 1220, 2610]
            elif num_cols == 6:
                col_widths_dxa = [2810, 1310, 1310, 1310, 1310, 1310]
            else:
                col_widths_dxa = [9360 // num_cols] * num_cols

            tbl = doc.add_table(rows=num_rows, cols=num_cols)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

            tblPr = tbl._tbl.tblPr
            tblBorders = OxmlElement("w:tblBorders")
            for side in ("top", "bottom", "insideH"):
                b = OxmlElement(f"w:{side}")
                b.set(qn("w:val"), "single")
                b.set(qn("w:sz"), "4")
                b.set(qn("w:space"), "0")
                b.set(qn("w:color"), "D3D3D3")
                tblBorders.append(b)
            for side in ("insideV", "left", "right"):
                b = OxmlElement(f"w:{side}")
                b.set(qn("w:val"), "none")
                tblBorders.append(b)
            tblPr.append(tblBorders)

            for r_idx, row_data in enumerate(data):
                row_cells = tbl.rows[r_idx].cells
                for c_idx in range(num_cols):
                    cell = row_cells[c_idx]
                    val = row_data[c_idx] if c_idx < len(row_data) else ""
                    set_cell_width_dxa(cell, col_widths_dxa[c_idx])
                    cp = cell.paragraphs[0]
                    if r_idx == 0:
                        set_cell_shading(cell, "003366")
                        set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
                        run = cp.add_run(val)
                        run.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        run.font.size = Pt(9.5)
                    else:
                        fill = "F9FBFD" if r_idx % 2 == 1 else "FFFFFF"
                        set_cell_shading(cell, fill)
                        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
                        run = cp.add_run(val)
                        run.font.size = Pt(9)

            spacer_after = doc.add_paragraph()
            spacer_after.paragraph_format.space_after = Pt(4)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def build_pdf(blocks, out_path: Path):
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=letter,
        topMargin=72,
        bottomMargin=72,
        leftMargin=72,
        rightMargin=72,
    )

    base = getSampleStyleSheet()
    st_inst = ParagraphStyle(
        "InstHeader",
        parent=base["Normal"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=15,
        textColor=colors.HexColor("#646464"),
        alignment=TA_CENTER,
        spaceBefore=24,
        spaceAfter=6,
    )
    st_title = ParagraphStyle(
        "DocTitle",
        parent=base["Normal"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#003366"),
        alignment=TA_CENTER,
        spaceAfter=18,
    )
    st_sub = ParagraphStyle(
        "DocSubtitle",
        parent=base["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#282828"),
        alignment=TA_CENTER,
        spaceAfter=24,
    )
    st_callout = ParagraphStyle(
        "Callout",
        parent=base["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=10.5,
        leading=14,
        textColor=colors.HexColor("#323232"),
    )
    st_h1 = ParagraphStyle(
        "H1",
        parent=base["Normal"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#003366"),
        spaceBefore=18,
        spaceAfter=6,
        keepWithNext=True,
    )
    st_h2 = ParagraphStyle(
        "H2",
        parent=base["Normal"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=17,
        textColor=colors.HexColor("#2980B9"),
        spaceBefore=14,
        spaceAfter=4,
        keepWithNext=True,
    )
    st_h3 = ParagraphStyle(
        "H3",
        parent=base["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#333333"),
        spaceBefore=10,
        spaceAfter=2,
        keepWithNext=True,
    )
    st_body = ParagraphStyle(
        "Body",
        parent=base["Normal"],
        fontName="Helvetica",
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#282828"),
        spaceBefore=0,
        spaceAfter=6,
    )
    st_bullet = ParagraphStyle(
        "BulletItem",
        parent=st_body,
        leftIndent=18,
        firstLineIndent=-10,
        spaceAfter=3,
    )
    st_th = ParagraphStyle(
        "TableHeader",
        parent=base["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9.5,
        leading=12,
        textColor=colors.white,
        alignment=TA_LEFT,
    )
    st_td = ParagraphStyle(
        "TableCell",
        parent=base["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=11.5,
        textColor=colors.HexColor("#282828"),
        alignment=TA_LEFT,
    )

    story = []
    usable_width = letter[0] - 144  # 468 pt (6.5 in)

    for btype, data in blocks:
        if btype == "inst_header":
            html = "<br/>".join(inline_md_to_rl_xml(line, "#646464") for line in data)
            story.append(Paragraph(html, st_inst))

        elif btype == "title":
            story.append(Paragraph(inline_md_to_rl_xml(data, "#003366"), st_title))

        elif btype == "subtitle":
            html = "<br/>".join(inline_md_to_rl_xml(line, "#282828") for line in data)
            story.append(Paragraph(html, st_sub))

        elif btype == "callout":
            title_line, body_text = data
            story.append(Spacer(1, 12))
            callout_html = ""
            if title_line:
                callout_html += (
                    f'<font name="Helvetica-Bold" size="11" color="#003366">'
                    f"{inline_md_to_rl_xml(title_line, '#003366')}</font><br/>"
                )
            callout_html += inline_md_to_rl_xml(body_text)
            t = Table([[Paragraph(callout_html, st_callout)]], colWidths=[usable_width])
            t.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F0F4F8")),
                        ("LINEBEFORE", (0, 0), (0, -1), 3, colors.HexColor("#003366")),
                        ("TOPPADDING", (0, 0), (-1, -1), 7),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                        ("LEFTPADDING", (0, 0), (-1, -1), 10),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                    ]
                )
            )
            story.append(t)
            story.append(Spacer(1, 4))

        elif btype == "h1":
            story.append(Paragraph(inline_md_to_rl_xml(data, "#003366"), st_h1))

        elif btype == "h2":
            story.append(Paragraph(inline_md_to_rl_xml(data, "#2980B9"), st_h2))

        elif btype == "h3":
            story.append(Paragraph(inline_md_to_rl_xml(data, "#333333"), st_h3))

        elif btype == "paragraph":
            story.append(Paragraph(inline_md_to_rl_xml(data), st_body))

        elif btype == "bullets":
            for item in data:
                story.append(Paragraph(f"•  {inline_md_to_rl_xml(item)}", st_bullet))

        elif btype == "table":
            num_cols = max(len(r) for r in data)
            if num_cols == 2:
                col_widths_pt = [usable_width * 0.26, usable_width * 0.74]
            elif num_cols == 5:
                col_widths_pt = [
                    usable_width * 0.22,
                    usable_width * 0.24,
                    usable_width * 0.13,
                    usable_width * 0.13,
                    usable_width * 0.28,
                ]
            elif num_cols == 6:
                col_widths_pt = [usable_width * 0.30] + [usable_width * 0.14] * 5
            else:
                col_widths_pt = [usable_width / num_cols] * num_cols
            rl_rows = []
            for r_idx, row_data in enumerate(data):
                row_cells = []
                for c_idx in range(num_cols):
                    val = row_data[c_idx] if c_idx < len(row_data) else ""
                    style = st_th if r_idx == 0 else st_td
                    row_cells.append(Paragraph(inline_md_to_rl_xml(val), style))
                rl_rows.append(row_cells)

            t = Table(rl_rows, colWidths=col_widths_pt, repeatRows=1)
            t_style = [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#003366")),
                ("TOPPADDING", (0, 0), (-1, 0), 6),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
                ("LEFTPADDING", (0, 0), (-1, 0), 6),
                ("RIGHTPADDING", (0, 0), (-1, 0), 6),
                ("TOPPADDING", (0, 1), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 1), (-1, -1), 5),
                ("LEFTPADDING", (0, 1), (-1, -1), 5),
                ("RIGHTPADDING", (0, 1), (-1, -1), 5),
                ("LINEABOVE", (0, 0), (-1, 0), 0.5, colors.HexColor("#D3D3D3")),
                ("LINEBELOW", (0, 0), (-1, -1), 0.5, colors.HexColor("#D3D3D3")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
            for r_idx in range(1, len(rl_rows)):
                bg = "#F9FBFD" if r_idx % 2 == 1 else "#FFFFFF"
                t_style.append(("BACKGROUND", (0, r_idx), (-1, r_idx), colors.HexColor(bg)))
            t.setStyle(TableStyle(t_style))
            story.append(t)
            story.append(Spacer(1, 6))

    doc.build(story)


def main():
    parser = argparse.ArgumentParser(
        description="Convierte el Markdown ground truth a DOCX y PDF con el formato de entrega."
    )
    parser.add_argument(
        "-i",
        "--input",
        default=str(REPO_ROOT / "Proyecto_Inversion_Injertos_Patagonia.md"),
        help="Ruta al archivo Markdown fuente.",
    )
    parser.add_argument(
        "--docx",
        default=str(REPO_ROOT / "entregas" / "Proyecto_Inversion_Injertos_Patagonia_Entregable.docx"),
        help="Ruta de salida para el archivo DOCX.",
    )
    parser.add_argument(
        "--pdf",
        default=str(REPO_ROOT / "entregas" / "Proyecto_Inversion_Injertos_Patagonia_Entregable.pdf"),
        help="Ruta de salida para el archivo PDF.",
    )
    args = parser.parse_args()

    md_path = Path(args.input)
    md_text = md_path.read_text(encoding="utf-8")
    blocks = parse_markdown_blocks(md_text)

    docx_path = Path(args.docx)
    pdf_path = Path(args.pdf)

    build_docx(blocks, docx_path)
    build_pdf(blocks, pdf_path)
    print(f"Generados:\n  DOCX: {docx_path}\n  PDF:  {pdf_path}")


if __name__ == "__main__":
    main()
